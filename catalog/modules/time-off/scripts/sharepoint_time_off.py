"""Downloads the team's time-off sheet from SharePoint and writes one shard per person.

The sheet is usually the results workbook of a "time-off request" form. Its
location is configuration, never code:

    TIMEOFF_SHAREPOINT_URL           view URL of the workbook (required online)
    TIMEOFF_SHAREPOINT_DOWNLOAD_URL  direct download URL, e.g.
                                     .../_layouts/15/download.aspx?UniqueId=...
                                     (required online)
    TIMEOFF_SHAREPOINT_META_URL      optional REST URL answering TimeLastModified
    TIMEOFF_SHAREPOINT_SITE_URL      optional lightweight page on the same site
                                     (defaults to the origin of the view URL)

Column names default to an English form export and can be overridden:

    TIMEOFF_SHEET           (default "Sheet1")
    TIMEOFF_COL_NAME        (default "Full name")      - the person who is OFF
    TIMEOFF_COL_EMAIL       (default "Email")          - the person who SUBMITTED the form
    TIMEOFF_COL_DATE_FROM   (default "First day off")
    TIMEOFF_COL_DATE_TO     (default "Last day off")
    TIMEOFF_COL_SUBMITTED   (default "Completion time")

Progress goes to stdout as single JSON lines {"step": ..., "message": ...},
ending with one {"status": "ok", ...} line. An error is an exception: its
text on stderr and exit code 1. server/index.js reads both streams and
forwards progress over SSE.

Offline mode (no session, no browser):
    python3 sharepoint_time_off.py --from-xlsx path.xlsx --dest .../time-off
"""

import argparse
import json
import os
import re
import sys
import time
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from _sso_shared import merge_file_into_seed, seed_exists  # noqa: E402


def _origin(url):
    try:
        parsed = urlparse(url or "")
    except ValueError:
        return ""
    if not parsed.scheme or not parsed.netloc:
        return ""
    return f"{parsed.scheme}://{parsed.netloc}/"


URL = os.environ.get("TIMEOFF_SHAREPOINT_URL", "").strip()
DOWNLOAD_URL = os.environ.get("TIMEOFF_SHAREPOINT_DOWNLOAD_URL", "").strip()
META_URL = os.environ.get("TIMEOFF_SHAREPOINT_META_URL", "").strip()
SITE_REFERER_URL = os.environ.get("TIMEOFF_SHAREPOINT_SITE_URL", "").strip() or _origin(URL)

LOGIN_HOST_FRAGMENTS = (
    "login.microsoftonline.com",
    "login.microsoft.com",
    "account.activedirectory.windowsazure.com",
    "login.live.com",
)

# Pages visited once after a login so every SharePoint host the download
# touches has its session cookie before the first REST call.
WARMUP_URLS = [u for u in dict.fromkeys([URL, SITE_REFERER_URL, _origin(DOWNLOAD_URL)]) if u]

SESSION_FRESH_HOURS = 168

SHEET_NAME = os.environ.get("TIMEOFF_SHEET", "Sheet1")
COL_NAME = os.environ.get("TIMEOFF_COL_NAME", "Full name")
COL_EMAIL = os.environ.get("TIMEOFF_COL_EMAIL", "Email")
COL_DATE_FROM = os.environ.get("TIMEOFF_COL_DATE_FROM", "First day off")
COL_DATE_TO = os.environ.get("TIMEOFF_COL_DATE_TO", "Last day off")
COL_SUBMITTED_AT = os.environ.get("TIMEOFF_COL_SUBMITTED", "Completion time")

PERSONS_SUBDIR = "people"

# Below this many known people the ratio guard stays quiet: a team of two
# losing one person is a real change, not a broken download.
MIN_PERSONS_FOR_RATIO_GUARD = 4
INDEX_SCHEMA = 1


def progress(step, message):
    """One JSON line on stdout - server/index.js forwards it over SSE."""
    print(json.dumps({"step": step, "message": message}, ensure_ascii=False), flush=True)


def require_config():
    """Online mode needs both URLs; say which one is missing."""
    missing = [name for name, value in (
        ("TIMEOFF_SHAREPOINT_URL", URL), ("TIMEOFF_SHAREPOINT_DOWNLOAD_URL", DOWNLOAD_URL),
    ) if not value]
    if missing:
        raise RuntimeError("Not configured: set " + " and ".join(missing) + " in .env.")


def _safe_close(browser):
    """Closes the browser without raising when it is already gone."""
    try:
        browser.close()
    except Exception:
        pass


def _looks_like_login(url):
    u = (url or "").lower()
    return any(frag in u for frag in LOGIN_HOST_FRAGMENTS)


def _host(url):
    try:
        return (urlparse(url or "").hostname or "").lower()
    except ValueError:
        return ""


# The hosts of the configured addresses - the only hosts that count as "the
# SharePoint site" (a tenant's personal OneDrive host and its site host
# differ, so both URLs contribute).
SITE_HOSTS = frozenset(h for h in (_host(URL), _host(DOWNLOAD_URL), _host(SITE_REFERER_URL), _host(META_URL)) if h)


def _is_sharepoint_url(url):
    """True only when the URL's HOST is one of the configured site hosts (not a substring anywhere)."""
    host = _host(url)
    return bool(host) and host in SITE_HOSTS


# -- Person slug ---------------------------------------------------------------
# Same rules as the app's name normaliser plus spaces -> hyphen. Letters with
# a stroke (l, d, o), which NFD does not decompose, go through an explicit map.
_STROKE_MAP = {
    "\u0142": "l", "\u0141": "l",
    "\u0111": "d", "\u0110": "d",
    "\u00f8": "o", "\u00d8": "o",
}
_TRAILING_APP_SCOPE_RE = re.compile(r"\s+app\s+(read|write|access)$")
_SEPARATOR_RUN_RE = re.compile(r"[-_\s]+")
_SPACE_RUN_RE = re.compile(r"\s+")


def slugify(name):
    """Turns a full name into the slug used as the shard file name."""
    s = str(name or "")
    s = unicodedata.normalize("NFD", s)
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    s = "".join(_STROKE_MAP.get(ch, ch) for ch in s)
    s = s.lower()
    s = _SEPARATOR_RUN_RE.sub(" ", s)
    # A name taken from a technical account sometimes carries a scope suffix,
    # e.g. "Jordan Lee App Write" must still match the person "Jordan Lee".
    s = _TRAILING_APP_SCOPE_RE.sub("", s)
    s = s.strip()
    s = _SPACE_RUN_RE.sub("-", s)
    return s


# -- SSO session -------------------------------------------------------------


def session_is_fresh(path, max_age_h=SESSION_FRESH_HOURS):
    try:
        age_h = (time.time() - path.stat().st_mtime) / 3600
    except OSError:
        return False
    return age_h < max_age_h


def session_valid(p, session_path):
    """Headless check: the session answers without a login redirect."""
    browser = p.chromium.launch(headless=True)
    try:
        context = browser.new_context(storage_state=str(session_path))
        page = context.new_page()
        page.goto(URL, wait_until="domcontentloaded", timeout=30_000)
        return not _looks_like_login(page.url)
    except Exception:
        return False
    finally:
        _safe_close(browser)


def _launch_headed(p):
    """Visible window in the person's own Chrome or Edge, Playwright's Chromium last.

    Interactive Microsoft sign-in with device-based Conditional Access passes
    in a real browser and can fail in "Chrome for Testing".
    """
    args = ["--ignore-certificate-errors"]
    for channel in ("chrome", "msedge"):
        try:
            return p.chromium.launch(channel=channel, headless=False, args=args)
        except Exception:  # noqa: BLE001 - channel not installed
            continue
    return p.chromium.launch(headless=False, args=args)


def silent_login(p, session_path, seed_path):
    """Headless sign-in from the shared Microsoft seed (_sso_shared.py), no window.

    True once the SharePoint pages answered without a login redirect and the
    session is saved; False otherwise (the caller then opens a window).
    """
    if not seed_exists(seed_path):
        return False
    progress("login", "Trying the saved Microsoft session (no window)...")
    browser = p.chromium.launch(headless=True)
    try:
        context = browser.new_context(storage_state=str(seed_path))
        page = context.new_page()
        for url in [u for u in (URL, SITE_REFERER_URL) if u]:
            try:
                page.goto(url, wait_until="domcontentloaded", timeout=25_000)
            except Exception:  # noqa: BLE001 - network or timeout
                return False
            current = page.url or ""
            if _looks_like_login(current) or not _is_sharepoint_url(current):
                return False
        context.storage_state(path=str(session_path))
        os.chmod(session_path, 0o600)
        return True
    except Exception:  # noqa: BLE001
        return False
    finally:
        _safe_close(browser)


def headed_login(p, session_path, timeout_s, seed_path=None):
    """Opens a visible browser window and waits for the person to sign in.

    The end of the sign-in is decided only by polling page.url once a second.
    """
    progress(
        "login",
        "Sign in in the browser window that just opened "
        "(Microsoft SSO + MFA, tick \"Stay signed in\")...",
    )
    browser = _launch_headed(p)
    context = browser.new_context(storage_state=str(seed_path)) if seed_exists(seed_path) else browser.new_context()
    page = context.new_page()
    try:
        page.goto(URL)
    except Exception:  # noqa: BLE001 - the login redirect chain may outlive goto's timeout
        pass

    deadline = time.monotonic() + timeout_s
    saved = False
    while time.monotonic() < deadline:
        page.wait_for_timeout(1000)
        try:
            current_url = page.url
        except Exception as exc:
            raise RuntimeError("The browser window was closed before signing in.") from exc
        if _is_sharepoint_url(current_url) and not _looks_like_login(current_url):
            context.storage_state(path=str(session_path))
            os.chmod(session_path, 0o600)
            saved = True
            break

    if not saved:
        _safe_close(browser)
        minutes = max(1, timeout_s // 60)
        raise RuntimeError(f"Not signed in within {minutes} minutes - try again.")

    _safe_close(browser)

    # Warm up every SharePoint host before the first download: the first
    # navigation on a site collection sets the FedAuth cookie the REST call
    # and the download itself need.
    browser2 = p.chromium.launch(headless=True)
    try:
        context2 = browser2.new_context(storage_state=str(session_path))
        any_ok = False
        for warm_url in WARMUP_URLS:
            try:
                warm_page = context2.new_page()
                warm_page.goto(warm_url, wait_until="domcontentloaded", timeout=20_000)
                warm_page.wait_for_timeout(1500)
                warm_page.close()
                any_ok = True
            except Exception:
                pass
        if any_ok:
            context2.storage_state(path=str(session_path))
            os.chmod(session_path, 0o600)
    finally:
        _safe_close(browser2)


def login_with_seed(p, session_path, timeout_s, seed_path=None):
    """Silent sign-in from the shared seed, else a visible window; merge the result back."""
    if not silent_login(p, session_path, seed_path):
        headed_login(p, session_path, timeout_s, seed_path)
    merge_file_into_seed(session_path, seed_path)


def ensure_session(p, session_path, timeout_s, seed_path=None):
    """Signs in again only when there is no session, or it is old or dead."""
    if session_path.exists() and session_is_fresh(session_path) and session_valid(p, session_path):
        progress("session", "The sign-in session is current.")
        return
    login_with_seed(p, session_path, timeout_s, seed_path)


KEEPALIVE_EXIT_EXPIRED = 4  # distinct from 1, which is any uncaught exception


def keepalive(session_path, seed_path=None):
    """Headless probe confirming the SharePoint session is still valid.

    Exit codes:
      0 - alive; session re-saved, prints {"status": "alive"}
      4 - not alive; SSO redirect, a non-2xx answer (401/403, error page)
          or an unreadable session file - the session file is never re-saved
      2 - unreachable; navigation failed or ended off the SharePoint host
          (VPN/network)
      anything else (1 from an uncaught exception, e.g. Playwright or its
      browser missing) - a failure, never reported as an expired session

    The caller (server/index.js) reads only the last JSON line from stdout
    and never sees cookie values - this function prints nothing else.
    """
    from playwright.sync_api import sync_playwright  # local import keeps Playwright optional

    if not session_path.exists():
        print(json.dumps({"status": "session-missing"}), flush=True)
        sys.exit(KEEPALIVE_EXIT_EXPIRED)

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        try:
            try:
                context = browser.new_context(storage_state=str(session_path))
            except Exception:  # noqa: BLE001 - unreadable session counts as expired
                sys.exit(KEEPALIVE_EXIT_EXPIRED)

            page = context.new_page()
            try:
                resp = page.goto(SITE_REFERER_URL, wait_until="domcontentloaded", timeout=15_000)
            except Exception:
                current = page.url or ""
                if _looks_like_login(current):
                    sys.exit(KEEPALIVE_EXIT_EXPIRED)
                sys.exit(2)

            if _looks_like_login(page.url):
                sys.exit(KEEPALIVE_EXIT_EXPIRED)

            # Off the SharePoint host (about:blank, a captive portal) is
            # unreachable; a 401/403 or any other error page is not alive.
            # Neither touches the session file, whose mtime is the freshness clock.
            if not _is_sharepoint_url(page.url):
                sys.exit(2)
            status = resp.status if resp is not None else None
            if status is None or not 200 <= status <= 299:
                sys.exit(KEEPALIVE_EXIT_EXPIRED)

            try:
                context.storage_state(path=str(session_path))
                try:
                    os.chmod(session_path, 0o600)
                except OSError:
                    pass
            except Exception:
                pass  # best-effort: a save failure must not kill the keepalive
        finally:
            _safe_close(browser)

    merge_file_into_seed(session_path, seed_path)
    print(json.dumps({"status": "alive"}), flush=True)


# -- Download and conversion -----------------------------------------------------


def download_xlsx(p, session_path, download_dir):
    """Downloads the workbook with the saved session into download_dir."""
    download_dir = Path(download_dir)
    download_dir.mkdir(parents=True, exist_ok=True)

    browser = p.chromium.launch(headless=True)
    try:
        context = browser.new_context(storage_state=str(session_path), accept_downloads=True)
        page = context.new_page()
        page.goto(URL, wait_until="domcontentloaded", timeout=30_000)

        if _looks_like_login(page.url):
            raise RuntimeError("The session expired during the download - start the download again.")

        with page.expect_download(timeout=30_000) as dl_info:
            try:
                page.goto(DOWNLOAD_URL, wait_until="commit", timeout=10_000)
            except Exception:
                pass  # a navigation that turns into a download always raises

        download = dl_info.value
        suggested = download.suggested_filename or "time-off.xlsx"
        dest_path = download_dir / suggested
        download.save_as(str(dest_path))
    finally:
        _safe_close(browser)

    return dest_path


def xlsx_to_records(path):
    """Workbook -> list of dicts keyed by the header row. Pure, no network."""
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    if not rows:
        return []

    headers = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
    records = []
    for row in rows[1:]:
        record = {}
        for header, cell in zip(headers, row):
            if isinstance(cell, datetime):
                record[header] = cell.isoformat()
            else:
                record[header] = cell
        records.append(record)
    return records


def fetch_remote_mtime(p, session_path):
    """TimeLastModified through REST. Best-effort: the caller catches the exception."""
    if not META_URL:
        raise RuntimeError("TIMEOFF_SHAREPOINT_META_URL is not set.")
    session_path = Path(session_path)
    if not session_path.exists():
        raise RuntimeError("No saved session - cannot read the modification date.")

    browser = p.chromium.launch(headless=True)
    try:
        context = browser.new_context(storage_state=str(session_path))
        page = context.new_page()
        page.goto(SITE_REFERER_URL, wait_until="domcontentloaded", timeout=15_000)
        if _looks_like_login(page.url):
            raise RuntimeError("The session expired while reading the modification date.")

        resp = context.request.get(
            META_URL,
            headers={"Accept": "application/json;odata=nometadata"},
            timeout=15_000,
        )
        body = resp.text()
        if not resp.ok:
            raise RuntimeError(f"Reading SharePoint metadata failed: HTTP {resp.status}")

        data = json.loads(body)
        mtime = data.get("TimeLastModified")
        if not mtime:
            raise RuntimeError("The SharePoint answer has no modification date.")
        return mtime
    finally:
        _safe_close(browser)


# -- Shards per person -----------------------------------------------------------


def records_to_shards(records):
    """Groups sheet records per person, slugged by the NAME column.

    NEVER slug by the e-mail column - that is whoever submitted the form, not
    the person who is off (one person often files time off for another). A
    record without a name or without both dates is skipped and counted.
    """
    shards = {}
    skipped = 0

    for record in records:
        raw_name = record.get(COL_NAME)
        name = str(raw_name).strip() if raw_name is not None else ""
        date_from = record.get(COL_DATE_FROM)
        date_to = record.get(COL_DATE_TO)

        if not name or not date_from or not date_to:
            skipped += 1
            continue

        slug = slugify(name)
        if not slug:
            skipped += 1
            continue

        submitted_at = record.get(COL_SUBMITTED_AT)
        raw_email = record.get(COL_EMAIL)
        submitted_by = str(raw_email).strip().lower() if raw_email else None

        entry = {
            "dateFrom": str(date_from)[:10],
            "dateTo": str(date_to)[:10],
            "submittedAt": str(submitted_at) if submitted_at is not None else None,
            "submittedBy": submitted_by,
        }

        shard = shards.setdefault(slug, {"name": name, "slug": slug, "entries": []})
        shard["entries"].append(entry)

    for shard in shards.values():
        shard["entries"].sort(key=lambda e: (e["dateFrom"], e["dateTo"], e["submittedAt"] or ""))

    return shards, skipped


def write_shards(dest, shards, fetched_at, source_mtime, skipped):
    """Writes one file per person (tmp+rename) and a combined index.json.

    A file is written only when its content really changed - two runs over
    the same data give changed == 0 and an untouched file mtime. A shard with
    no entries in this run is removed (a person removed from the sheet goes
    away with the next push - intended).
    """
    dest = Path(dest)
    persons_dir = dest / PERSONS_SUBDIR
    persons_dir.mkdir(parents=True, exist_ok=True)

    index_path = dest / "index.json"
    prev_persons = {}
    if index_path.exists():
        try:
            prev_persons = json.loads(index_path.read_text(encoding="utf-8")).get("persons", {}) or {}
        except (json.JSONDecodeError, OSError):
            prev_persons = {}

    # Destructive-ratio guard: the removal loop below deletes every person
    # file this run did not produce, so an empty or half-downloaded sheet
    # would erase the whole team's history in one pass. A live form-results
    # sheet only ever grows, so losing more than half of the known people is
    # a broken download, never a real change - refuse before touching a file.
    existing_files = sorted(persons_dir.glob("*.json"))
    if existing_files:
        vanishing = [f for f in existing_files if f.stem not in shards]
        if not shards:
            raise RuntimeError(
                "The sheet has no valid entry, but there is data for "
                f"{len(existing_files)} people on disk - not deleting it. Check the downloaded file."
            )
        if len(existing_files) >= MIN_PERSONS_FOR_RATIO_GUARD and len(vanishing) * 2 > len(existing_files):
            raise RuntimeError(
                f"{len(vanishing)} of {len(existing_files)} people vanished from the sheet - that looks "
                "like a broken or partial file, so nothing is deleted. Check the downloaded file."
            )

    changed = 0
    persons_out = {}
    keep_files = set()

    for slug, shard in shards.items():
        serialized = json.dumps(shard, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
        file_path = persons_dir / f"{slug}.json"
        keep_files.add(file_path.name)

        existing = None
        if file_path.exists():
            try:
                existing = file_path.read_text(encoding="utf-8")
            except OSError:
                existing = None

        if existing == serialized:
            last_changed = (prev_persons.get(slug) or {}).get("last_changed", fetched_at)
        else:
            tmp_path = persons_dir / f"{slug}.json.tmp"
            tmp_path.write_text(serialized, encoding="utf-8")
            os.replace(str(tmp_path), str(file_path))
            last_changed = fetched_at
            changed += 1

        persons_out[slug] = {
            "name": shard["name"],
            "last_changed": last_changed,
            "count": len(shard["entries"]),
        }

    removed = 0
    for existing_file in persons_dir.glob("*.json"):
        if existing_file.name not in keep_files:
            existing_file.unlink()
            removed += 1

    index_payload = {
        "schema": INDEX_SCHEMA,
        "fetched_at": fetched_at,
        "source_mtime": source_mtime,
        "count": sum(len(s["entries"]) for s in shards.values()),
        "skipped": skipped,
        "persons": persons_out,
    }
    index_serialized = json.dumps(index_payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n"
    index_tmp = dest / "index.json.tmp"
    index_tmp.write_text(index_serialized, encoding="utf-8")
    os.replace(str(index_tmp), str(index_path))

    return {
        "count": index_payload["count"],
        "persons": len(shards),
        "changed": changed,
        "removed": removed,
        "skipped": skipped,
    }


# -- CLI -------------------------------------------------------------------------


def parse_args(argv=None):
    parser = argparse.ArgumentParser(description="Downloads the team's time-off sheet from SharePoint.")
    parser.add_argument("--dest", default=None, help="The time-off/ directory in the team data, where shards go.")
    parser.add_argument("--session", default=None,
                        help="SSO session file (Playwright storage_state). Required outside --from-xlsx.")
    parser.add_argument("--download-dir", default=None,
                        help="Working directory for the downloaded XLSX. Required outside --from-xlsx and --keepalive.")
    parser.add_argument(
        "--headed-timeout", type=int, default=300,
        help="Seconds allowed for signing in in the opened browser window.",
    )
    parser.add_argument(
        "--from-xlsx", default=None,
        help="Offline mode: convert the given XLSX instead of signing in and downloading.",
    )
    parser.add_argument(
        "--seed", default=None,
        help="Shared Microsoft session of every SSO module (_sso_shared.py).",
    )
    parser.add_argument(
        "--login", action="store_true",
        help="Sign in only: opens a visible SSO window and saves the session; needs --session; exit 0 signed in, exit 1 not.",
    )
    parser.add_argument(
        "--keepalive", action="store_true",
        help="Headless probe confirming the SSO session; needs --session; exit 0 alive, exit 4 expired, exit 2 no network.",
    )
    args = parser.parse_args(argv)
    if args.login:
        if not args.session:
            parser.error("--login needs --session <file>.")
    elif args.keepalive:
        if not args.session:
            parser.error("--keepalive needs --session <file>.")
    elif not args.from_xlsx and (not args.session or not args.download_dir or not args.dest):
        parser.error("--session, --download-dir and --dest are required outside --from-xlsx and --keepalive.")
    elif args.from_xlsx and not args.dest:
        parser.error("--dest is required with --from-xlsx.")
    return args


def main():
    args = parse_args()

    seed_path = Path(args.seed) if args.seed else None

    if args.login:
        require_config()
        from playwright.sync_api import sync_playwright  # local import keeps Playwright optional
        with sync_playwright() as p:
            try:
                login_with_seed(p, Path(args.session), args.headed_timeout, seed_path)
            except RuntimeError as exc:
                progress("login", str(exc))
                sys.exit(1)
        print(json.dumps({"status": "logged-in"}), flush=True)
        return

    if args.keepalive:
        require_config()
        keepalive(Path(args.session), seed_path)
        return

    dest = Path(args.dest)
    source_mtime = None

    if args.from_xlsx:
        xlsx_path = Path(args.from_xlsx)
        records = xlsx_to_records(xlsx_path)
        progress("convert", f"Processed {len(records)} sheet rows (offline mode).")
    else:
        require_config()
        from playwright.sync_api import sync_playwright  # local import keeps Playwright optional

        session_path = Path(args.session)
        download_dir = Path(args.download_dir)

        progress("session", "Checking the saved sign-in session...")
        with sync_playwright() as p:
            ensure_session(p, session_path, args.headed_timeout, seed_path)

            progress("download", "Downloading the sheet from SharePoint...")
            xlsx_path = download_xlsx(p, session_path, download_dir)

            if META_URL:
                try:
                    source_mtime = fetch_remote_mtime(p, session_path)
                except Exception:
                    progress("meta", "Could not read the sheet's modification date - saving without it.")
                    source_mtime = None

        records = xlsx_to_records(xlsx_path)
        progress("convert", f"Processed {len(records)} sheet rows.")
        try:
            xlsx_path.unlink()
        except OSError:
            pass

    shards, skipped = records_to_shards(records)
    fetched_at = datetime.now(timezone.utc).isoformat()

    result = write_shards(dest, shards, fetched_at, source_mtime, skipped)
    progress(
        "write",
        f"Saving data: {result['persons']} people, {result['changed']} changed files...",
    )

    print(json.dumps({
        "status": "ok",
        "count": result["count"],
        "persons": result["persons"],
        "changed": result["changed"],
        "removed": result["removed"],
        "skipped": result["skipped"],
        "fetched_at": fetched_at,
        "source_mtime": source_mtime,
    }, ensure_ascii=False))


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(str(exc), file=sys.stderr)
        sys.exit(1)
